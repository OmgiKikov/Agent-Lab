#!/usr/bin/env python3
"""Prepare the original conversation rows for local LangWatch import."""

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: prepare-dialogues.py SOURCE.xlsx OUTPUT.jsonl")

    source = Path(sys.argv[1])
    destination = Path(sys.argv[2])
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["Данные"]
    source_rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(source_rows)]
    positions = {name: index for index, name in enumerate(headers)}
    required = ("Id диалога", "Дата", "dialogSummary", "Канал", "Источник", "Текст")
    missing = [name for name in required if name not in positions]
    if missing:
        raise SystemExit(f"Missing source columns: {', '.join(missing)}")

    if destination.suffix != ".jsonl":
        raise SystemExit("The output must be JSONL so Cyrillic text survives the LangWatch upload")

    destination.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with destination.open("w", encoding="utf-8", newline="") as output:
        for row_number, row in enumerate(source_rows, start=2):
            if all(value is None for value in row):
                continue
            original = {name: row[index] for index, name in enumerate(headers)}
            entry = {
                "source_row": str(row_number),
                "dialogue_id": str(original["Id диалога"] or ""),
                "conversation": str(original["Текст"] or ""),
                "date": str(original["Дата"] or ""),
                "channel": str(original["Канал"] or ""),
                "source": str(original["Источник"] or ""),
                "summary": str(original["dialogSummary"] or ""),
                "source_metadata": json.dumps(
                    {name: value for name, value in original.items() if name != "Текст"},
                    ensure_ascii=False,
                    default=str,
                ),
            }
            output.write(json.dumps(entry, ensure_ascii=True) + "\n")
            count += 1
    print(f"Prepared {count} conversations in {destination}")


if __name__ == "__main__":
    main()

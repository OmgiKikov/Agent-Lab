#!/usr/bin/env python3
"""Разбор сломавшихся кейсов agent_oc (.xlsx) → задание end2end для Agent Lab.

Файл разбора содержит всё, чего не хватает карточке: дословный вопрос клиента, условия кейса
(полномочия, профиль, ЕПК), поверхность и ожидание владельца в колонке «Ожидалось». Поэтому
карточки собираются из него, а не пишутся руками и не придумываются моделью.

    python examples/agent-oc-cases.py --input "сломавшиеся кейсы раг.xlsx" --output cases-task.json
    AGENT_LAB_TASK=cases-task.json AGENT_OC_EPK_UL=… AGENT_OC_EPK_FL=… bash examples/agent-oc-e2e.sh

По умолчанию берутся строки, где «Версия» = «Новая», а «Итог проверки» = «Ошибка», то есть то,
что сломалось на текущей редакции. `--passing` берёт вместо этого успешные строки — тогда прогон
проверяет, что рабочее не развалилось.

Текстовое ожидание владельца становится `successCriteria` и рубрикой судьи дословно. Ожидание,
записанное голым кодом ответа («202-2»), становится точной проверкой `result.status_code`: код
кладёт туда адаптер, а судья смысла кодов агента не знает.

Нужен `openpyxl`. Скрипт ничего не импортирует из agent_oc.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional, Tuple

SURFACES = {
    "textassistant": "TEXTASSISTANT",
    "gigaassistant": "GIGAASSISTANT",
    "voiceassistant": "VOICEASSISTANT",
    "voicecopilot": "VOICECOPILOT",
    "textcopilot": "TEXTCOPILOT",
    "ocassistant": "OCASSISTANT",
}
_AUTHORITY = re.compile(r"Полномочия\s*:\s*(\d)")
_EPK_UL = re.compile(r"EPC\s*UL\s*:\s*(\d+)", re.IGNORECASE)
_EPK_FL = re.compile(r"EPC\s*FL\s*:\s*(\d+)", re.IGNORECASE)
_SLUG = re.compile(r"[^a-zA-Zа-яА-Я0-9]+")
# Часть ожиданий в разборе записана голым кодом ответа («202-2», «202-1»). Судья не знает, что
# означают коды агента, и в живом прогоне засчитал инструкцию по заказу инкассации как «202-1».
# Код ответа адаптер кладёт в result.status_code, поэтому такое ожидание проверяется объективно.
_BARE_CODE = re.compile(r"^\s*(\d{3}(?:-\d+)?)\s*$")

MAX_CRITERIA = 3000
MAX_RUBRIC = 2000
MAX_CASES = 40


def slug(text: str, index: int) -> str:
    words = [word for word in _SLUG.split(text.lower()) if word][:4]
    latin = "-".join(re.sub(r"[^a-z0-9]", "", transliterate(word)) for word in words)
    return f"{latin or 'case'}-{index}"[:80]


TRANSLIT = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh", "з": "z", "и": "i",
    "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t",
    "у": "u", "ф": "f", "х": "h", "ц": "c", "ч": "ch", "ш": "sh", "щ": "sch", "ъ": "", "ы": "y", "ь": "",
    "э": "e", "ю": "yu", "я": "ya",
}


def transliterate(word: str) -> str:
    return "".join(TRANSLIT.get(letter, letter) for letter in word)


def read_rows(path: Path, sheets: Optional[List[str]]) -> Iterator[Tuple[str, Dict[str, Any]]]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.sheetnames:
        if sheets and sheet not in sheets:
            continue
        rows = workbook[sheet].iter_rows(values_only=True)
        header = [str(cell or "") for cell in next(rows)]
        for row in rows:
            yield sheet, dict(zip(header, row))


def session(conditions: str, surface: str, epk_ul: str, epk_fl: str) -> Dict[str, Any]:
    authority = _AUTHORITY.search(conditions)
    ul = _EPK_UL.search(conditions)
    fl = _EPK_FL.search(conditions)
    return {
        "surface": SURFACES.get(surface.strip().lower(), surface.strip().upper()),
        "authority": authority.group(1) if authority else "3",
        "epk_ul": ul.group(1) if ul else epk_ul,
        "epk_fl": fl.group(1) if fl else epk_fl,
    }


def case(row: Dict[str, Any], index: int, epk_ul: str, epk_fl: str) -> Dict[str, Any]:
    question = str(row.get("Вопрос клиента") or "").strip()
    expected = str(row.get("Ожидалось") or "").strip()
    conditions = str(row.get("Условия кейса") or "")
    identifier = slug(question, index)
    return {
        "id": identifier,
        "title": question[:200],
        "goal": f"Получить от агента ответ на вопрос: «{question}»"[:3000],
        "opening": question[:3000],
        "facts": "Клиент знает только то, что написал в своём вопросе.",
        "behavior": "Отвечай на уточнения агента только известными фактами; закончи, когда получишь ответ по существу или поймёшь, что его не будет.",
        "maxFollowUps": 2,
        "initialState": {
            "records": {
                "session": session(conditions, str(row.get("Поверхность") or ""), epk_ul, epk_fl),
                "result": {"status_code": "", "produced_by": "", "seconds": 0},
            },
            "writableFields": ["status_code", "produced_by", "seconds"],
            "transientFailures": 0,
        },
        **criteria(expected),
    }


def criteria(expected: str) -> Dict[str, Any]:
    """Как проверять ожидание: голый код — точной проверкой состояния, текст — рубрикой судьи."""
    code = _BARE_CODE.match(expected)
    if code:
        return {
            "successCriteria": f"Агент завершает диалог кодом ответа {code.group(1)}",
            "checks": [{
                "id": "status_code", "kind": "state_equals", "stage": "understand",
                "description": f"Код ответа агента {code.group(1)}, как в разборе кейсов",
                "recordId": "result", "field": "status_code", "value": code.group(1),
            }],
            "metrics": [],
        }
    return {
        "successCriteria": f"Ответ агента по смыслу совпадает с ожиданием владельца: {expected}"[:MAX_CRITERIA],
        "checks": [],
        "metrics": [{
            "id": "expected_answer",
            "name": "Совпадение с ожиданием владельца",
            "subject": "agent",
            "stage": "compose",
            "description": f"Ожидание из разбора кейсов: {expected}"[:MAX_RUBRIC],
            "passCriteria": f"Ответ агента передаёт то же по существу, что ожидание владельца: {expected}"[:MAX_RUBRIC],
            "failCriteria": "Агент отвечает о другом, отказывается по теме, уводит в другой сценарий или не доходит до ответа."[:MAX_RUBRIC],
        }],
    }


def build(args: argparse.Namespace) -> Dict[str, Any]:
    wanted = "Успешно" if args.passing else "Ошибка"
    cases: List[Dict[str, Any]] = []
    seen: set[str] = set()
    for _, row in read_rows(Path(args.input), args.sheet):
        if str(row.get("Версия") or "").strip() != args.version:
            continue
        if str(row.get("Итог проверки") or "").strip() != wanted:
            continue
        if not str(row.get("Вопрос клиента") or "").strip() or not str(row.get("Ожидалось") or "").strip():
            continue
        item = case(row, len(cases) + 1, args.epk_ul, args.epk_fl)
        if item["id"] in seen:
            continue
        seen.add(item["id"])
        cases.append(item)
        if len(cases) >= args.limit:
            break

    if not cases:
        raise SystemExit("Под фильтры не попал ни один кейс: проверьте колонки «Версия» и «Итог проверки».")

    return {
        "task": f"Проверить агента инкассации на кейсах из разбора: {'что работало' if args.passing else 'что сломалось'} на редакции «{args.version}»",
        "mode": "live",
        "scenarioCount": 0,
        "materials": [{
            "name": "Разбор кейсов: вопрос клиента и ожидание владельца",
            "content": "\n\n".join(f"Вопрос: {c['opening']}\nОжидание владельца: {c['successCriteria']}" for c in cases)[:120000],
        }],
        "settings": {
            "provider": "giga", "model": args.model,
            "roles": {"judge": {"provider": "giga", "model": args.judge}},
            "userModes": ["reactive"], "repeats": 1, "maxTurns": 6,
            # Лимиты по верхней границе Agent Lab: на карточку уходит диалог плюс по два вызова
            # судьи на каждую применимую рубрику, и при десятиминутном умолчании прогон на
            # десятке карточек обрывается уже после диалогов — вердиктов не остаётся вовсе.
            "maxDurationMs": 3600000, "maxCalls": 3000,
        },
        "goldenCases": cases,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", required=True, help="Файл разбора кейсов .xlsx")
    parser.add_argument("--output", required=True, help="Куда записать задание для прогона")
    parser.add_argument("--sheet", action="append", help="Лист; можно повторять. По умолчанию все.")
    parser.add_argument("--version", default="Новая", help="Значение колонки «Версия» (по умолчанию «Новая»)")
    parser.add_argument("--passing", action="store_true", help="Взять успешные кейсы вместо сломавшихся")
    parser.add_argument("--limit", type=int, default=MAX_CASES, help=f"Не больше {MAX_CASES} карточек (предел Agent Lab); на первый прогон разумно 5-8")
    parser.add_argument("--epk-ul", default="EPK_UL_PLACEHOLDER", help="ЕПК ЮЛ, если его нет в условиях кейса")
    parser.add_argument("--epk-fl", default="EPK_FL_PLACEHOLDER", help="ЕПК ФЛ, если его нет в условиях кейса")
    parser.add_argument("--model", default="glm-5.2", help="Модель симулятора")
    parser.add_argument("--judge", default="GigaChat-3-Ultra", help="Модель судьи")
    parser.add_argument("--criteria-output", help="Файл критериев для переоценки уже записанного прогона (reassess)")
    parser.add_argument("--golden-output", help="Те же карточки массивом golden-кейсов для разговора (agent_lab_build goldenFile)")
    args = parser.parse_args()

    if args.limit > MAX_CASES:
        parser.error(f"Agent Lab принимает не больше {MAX_CASES} карточек за прогон.")

    task = build(args)
    Path(args.output).write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if args.golden_output:
        # В разговоре карточки передаются инструменту agent_lab_build файлом golden-кейсов: это та же
        # карточка, что в задании, без обвязки задания (моделей, лимитов, материалов).
        Path(args.golden_output).write_text(json.dumps(task["goldenCases"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if args.criteria_output:
        # Переоценка меняет критерии только у перечисленных карточек, а точные проверки считает по
        # сохранённым фактам - агент заново не вызывается.
        patch = {"criteria": [{"scenarioId": c["id"], "successCriteria": c["successCriteria"], "checks": c["checks"], "metrics": c["metrics"]}
                              for c in task["goldenCases"]]}
        Path(args.criteria_output).write_text(json.dumps(patch, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    surfaces = sorted({c["initialState"]["records"]["session"]["surface"] for c in task["goldenCases"]})
    count = len(task["goldenCases"])
    objective = sum(1 for c in task["goldenCases"] if c["checks"])
    print(f"{args.output}: карточек {count} (с точной проверкой кода ответа: {objective}), поверхности: {', '.join(surfaces)}")
    print(f"Ожидайте примерно {count} диалогов и до {count * 4} вызовов судьи; лимит прогона - час.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

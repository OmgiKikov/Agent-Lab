"""Logged conversations of the chat: the Excel export uploaded on the page becomes data/logs.jsonl.

The export (sheet «Данные») has one row per conversation: «Id диалога», and «Текст» with the turns marked
CLIENT and AGENT. It repeats every exchange; «Порядок сообщения в диалоге» lists the real messages.
"""

import io
import json
import re

from openpyxl import load_workbook

from .settings import DATA

FILE = DATA / 'logs.jsonl'
SHEET = 'Данные'
ID, TEXT, ORDER = 'Id диалога', 'Текст', 'Порядок сообщения в диалоге'
MARKER = re.compile(r'\b(CLIENT|AGENT)\b')


def load() -> list[dict]:
    """Conversations that start with the customer and have the agent's answer."""
    if not FILE.exists():
        return []
    rows = [json.loads(line) for line in FILE.read_text(encoding='utf-8').splitlines() if line.strip()]
    return [r for r in rows if len(r.get('messages') or []) >= 2 and r['messages'][0].get('role') == 'user']


def turns(text: str) -> list[dict]:
    marks = list(MARKER.finditer(text))
    return [
        {
            'role': 'user' if mark.group(1) == 'CLIENT' else 'assistant',
            'content': text[mark.end() : marks[i + 1].start() if i + 1 < len(marks) else len(text)].strip(),
        }
        for i, mark in enumerate(marks)
    ]


def collapse(messages: list[dict]) -> list[dict]:
    """Drop an exchange (customer + agent) that repeats the one right before it."""
    out: list[dict] = []
    for i in range(0, len(messages), 2):
        pair = messages[i : i + 2]
        if out[-2:] != pair:
            out += pair
    return out


def _count(order: object) -> int | None:
    try:
        return len(json.loads(str(order)))
    except (ValueError, TypeError):
        return None


def from_excel(data: bytes) -> list[dict]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook[SHEET] if SHEET in workbook.sheetnames else workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(value or '').strip() for value in next(rows, ())]
    if ID not in header or TEXT not in header:
        raise ValueError(f'В выгрузке нет колонок «{ID}» и «{TEXT}» (лист «{SHEET}»)')
    column = {name: header.index(name) for name in (ID, TEXT, ORDER) if name in header}
    dialogues = []
    for row in rows:
        text = str(row[column[TEXT]] or '')
        messages = collapse(turns(text))
        count = _count(row[column[ORDER]]) if ORDER in column else None
        if count:
            messages = messages[:count]
        dialogues.append({'id': str(row[column[ID]] or ''), 'messages': messages})
    return dialogues


def from_jsonl(data: bytes) -> list[dict]:
    return [json.loads(line) for line in data.decode('utf-8').splitlines() if line.strip()]


def replace(name: str, data: bytes) -> int:
    """Replace the logs with an uploaded export (.xlsx) or prepared conversations (.jsonl); how many are usable."""
    dialogues = from_jsonl(data) if name.lower().endswith('.jsonl') else from_excel(data)
    usable = [d for d in dialogues if d['id'] and len(d['messages']) >= 2 and d['messages'][0]['role'] == 'user']
    if not usable:
        raise ValueError('В файле нет разговоров, которые начинаются с клиента и содержат ответ агента')
    FILE.parent.mkdir(parents=True, exist_ok=True)
    temp = FILE.with_suffix('.tmp')
    temp.write_text(''.join(json.dumps(d, ensure_ascii=False) + '\n' for d in usable), encoding='utf-8')
    temp.chmod(0o600)
    temp.replace(FILE)
    return len(usable)

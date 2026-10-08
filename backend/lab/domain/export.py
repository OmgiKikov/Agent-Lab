"""An export of conversations from the chat (a workbook or JSON lines) read into conversations, and the text of a
conversation as the customer saw it, with the agent's own words in it.

Excel exports contain CLIENT/AGENT turns, on lines of their own or all on one line, and an order column. Exports
repeat exchanges: every exchange twice, or only some of them, sometimes with a trailing dot on the copy. A run of
identical adjacent exchanges is read as one or more real exchanges, and the repeats are removed only when the order
column's count admits exactly one reading; a customer's actual repeated question survives whenever the count requires
it. Otherwise the text is the conversation: the bank's export often counts other messages than its text holds, and a
row the count cannot settle is read as written, never refused or cut, with the reason in its meta.

Each conversation keeps the export's service columns (`meta`): the date, channel, entry point, the agents that
took part, each status code with the agents it names, and the operator flag. Customer identifiers are kept only as a
pseudonymous key for grouping. Parsing is pure so a cancelled import cannot commit from a thread.
"""

import ast
import hashlib
import io
import json
import re
import zlib
from xml.etree.ElementTree import ParseError
from zipfile import ZIP_DEFLATED, ZIP_STORED, BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from . import export_formats
from .transcript import calls_line

LIMIT = 50_000_000  # an uploaded export
INFLATED = 500_000_000  # the parts of a workbook, unpacked together
SHEET = 'Данные'
ID, TEXT, ORDER = 'Id диалога', 'Текст', 'Порядок сообщения в диалоге'
# Service columns read into meta: column → key. Customer ids are hashed into clientKey, never stored as is.
COLUMNS = {
    'Дата': 'date',
    'dialogSummary': 'summary',
    'agentCode': 'agents',
    'actionCode': 'actions',
    'Канал': 'channel',
    'Источник': 'entrySource',
    'Поверхность': 'surface',
    'Вызов оператора в диалоге': 'operator',
}
STATUS_COLUMNS = ('Статус код 200', 'Статус код 202_1', 'Статус код 202_2', 'Статус код 202_5', 'Статус код 202_7')
CLIENT_COLUMNS = ('epkId', 'sflEpkId', 'digitalUserId')
MAX_READINGS = 2  # counting stops here: two readings already make a conversation ambiguous
# An export that starts every turn on its own line: «HOST AGENT NOT FOUND» inside a message is the customer's words.
MARKER = re.compile(r'^[ \t]*(CLIENT|AGENT)\b', re.M)
# The bank's export writes a whole conversation on one line: «CLIENT … AGENT … CLIENT …». There a marker word inside a
# message cannot be told from a turn: the order column decides between the layouts, or the text is read as written.
INLINE = re.compile(r'(?:^|(?<=\s))(CLIENT|AGENT)\b')  # «CLIENT?»: the customer sent a lone «?»
# A broken workbook: openpyxl names a part the archive does not have (KeyError), zipfile meets a broken stream or a
# feature it does not read.
UNREADABLE = (BadZipFile, InvalidFileException, ParseError, KeyError, zlib.error, EOFError, NotImplementedError)
# A chat button the agent sent, written into the export's text as «` ` ` transition-code CODE ` ` `».
CONTROL = re.compile(r'`\s*`\s*`\s*transition-code\s*([A-Za-z0-9_-]*)\s*`\s*`\s*`')
# The lines the Lab puts under a reply: the buttons it sent (as_seen) and the systems it called (conversation).
LAB_LINES = re.compile(r'(?:\n\[(?:Кнопки|вызовы систем): [^\n]*\])+\Z')


def as_seen(text: str) -> str:
    """An agent reply as the customer saw it: its words, then the buttons it sent, not the export's control code.
    A judge shown the raw code reads it as the agent's formatting (frontend/src/product/text.ts does the same)."""
    buttons = [code or 'кнопка' for code in CONTROL.findall(text)]
    words = CONTROL.sub('', text).strip()
    return words + ('\n[Кнопки: ' + ' | '.join(buttons) + ']' if buttons else '')


def words(text: str) -> str:
    """What the agent wrote in a reply, raw or as the judge reads it: the lines of buttons and of the systems called are
    the Lab's and the code is the export's, so none of them is evidence of the agent's words."""
    return LAB_LINES.sub('', CONTROL.sub('', text)).strip()


def conversation(dialogue: dict) -> list[dict]:
    """A logged conversation as the judge reads it: the customer's words, and the agent's replies as the customer saw
    them (as_seen). A reply the agent has just given to recorded questions also names the systems it called, as in a
    played conversation (transcript.calls_line): an export records none."""
    return [
        {'role': 'CUSTOMER', 'text': m['content']}
        if m['role'] == 'user'
        else {'role': 'AGENT', 'text': as_seen(m['content']) + calls_line(m)}
        for m in dialogue['messages']
    ]


def preview(dialogue: dict) -> dict:
    """A conversation in a list: its id, the start of its first message and how many messages it has."""
    messages = dialogue['messages']
    return {'id': dialogue['id'], 'opening': messages[0]['content'][:240] if messages else '', 'turns': len(messages)}


def turns(text: str) -> list[dict]:
    """A conversation's text read as written, where no order column can settle its repeats (a .csv export)."""
    return _as_written(layouts(text))[1]


def _as_written(found: list[tuple[str, list[dict]]]) -> tuple[str, list[dict]]:
    """The layout a text is read in when no order column settles it: the likelier one that holds both sides of the
    conversation, else the likelier. Inline exchanges on lines of their own («CLIENT … AGENT …» per line) read on lines
    would put the agent's words into the customer's messages."""
    return next((f for f in found if {turn['role'] for turn in f[1]} == {'user', 'assistant'}), found[0])


def _split(text: str, marks: list[re.Match]) -> list[dict]:
    return [
        {
            'role': 'user' if mark.group(1) == 'CLIENT' else 'assistant',
            'content': text[mark.end() : marks[i + 1].start() if i + 1 < len(marks) else len(text)].strip(),
        }
        for i, mark in enumerate(marks)
    ]


def layouts(text: str) -> list[tuple[str, list[dict]]]:
    """The ways the exported text can be read into turns, the likelier first: turns on lines of their own, then turns
    written inline. A text can mix them (inline exchanges with line breaks between them or inside a message), so both
    are tried and the order column decides (_read_turns). A marker with nothing after it is no message."""
    lines, inline = list(MARKER.finditer(text)), list(INLINE.finditer(text))
    found = [('lines', _split(text, lines))] if len(lines) >= 2 else []
    if len(inline) > len(lines) or not found:
        found.append(('inline', _split(text, inline)))
    return [(layout, [turn for turn in turns if turn['content']]) for layout, turns in found]


class Unsettled(ValueError):
    """The order column cannot settle how this text reads; the text is then read as written."""


def _read_turns(text: str, order: object) -> tuple[str, list[dict], list[dict], list[int], str | None]:
    """The layout that reads the text into the order column's count of messages: (layout, exported turns, real
    messages, their positions, None). When none does, the likelier layout as written, and why the count did not
    settle it."""
    found = layouts(text)
    reason: Unsettled | None = None
    try:
        count = _message_count(order)
    except Unsettled as error:
        reason = error
    else:
        for layout, turns in found:
            try:
                messages, kept = _export_messages(turns, count)
            except Unsettled as error:
                reason = reason or error
                continue
            return layout, turns, messages, kept, None
    layout, turns = _as_written(found)
    return layout, turns, turns, list(range(len(turns))), str(reason or 'в тексте нет реплик')


def _message_count(order: object) -> int:
    """How many messages the order column lists; Unsettled when it lists none or cannot be read."""
    try:
        value = json.loads(str(order))
    except (ValueError, TypeError) as error:
        raise Unsettled('не читается порядок сообщений') from error
    if not isinstance(value, list) or not value:
        raise Unsettled('пустой порядок сообщений')
    return len(value)


def _same(text: str) -> str:
    """Comparison form of an exported message: the copy differs only in spacing and a trailing dot."""
    return re.sub(r'\s+', ' ', text).strip().rstrip('. ').strip()


def _exchanges(messages: list[dict]) -> list[list[int]]:
    """Message indexes grouped into exchanges: a customer message and the agent messages after it."""
    groups: list[list[int]] = []
    for index, message in enumerate(messages):
        if message['role'] == 'user' or not groups:
            groups.append([index])
        else:
            groups[-1].append(index)
    return groups


def _readings(sizes: list[int], repeats: list[int], target: int) -> tuple[int, list[int]]:
    """How many choices of a real multiplicity 1..r for each run of r identical exchanges reach the target count
    (capped at MAX_READINGS), and the choice itself when it is the only one."""
    ways = [[0] * (target + 1) for _ in range(len(sizes) + 1)]
    ways[len(sizes)][0] = 1
    for run in range(len(sizes) - 1, -1, -1):
        for total in range(target + 1):
            rests = (total - m * sizes[run] for m in range(1, repeats[run] + 1))
            ways[run][total] = min(MAX_READINGS, sum(ways[run + 1][rest] for rest in rests if rest >= 0))
    if ways[0][target] != 1:
        return ways[0][target], []
    chosen, total = [], target
    for run in range(len(sizes)):
        multiplicity = next(
            m
            for m in range(1, repeats[run] + 1)
            if total - m * sizes[run] >= 0 and ways[run + 1][total - m * sizes[run]]
        )
        chosen.append(multiplicity)
        total -= multiplicity * sizes[run]
    return 1, chosen


def _export_messages(messages: list[dict], count: int) -> tuple[list[dict], list[int]]:
    """The real messages and their positions in the exported text; Unsettled when the count cannot decide."""
    if len(messages) == count:
        return messages, list(range(len(messages)))
    # Known doubled layout: every exchange followed by its exact copy. It decides even where several runs could trade
    # multiplicities, e.g. a real repeated exchange exported twice.
    if count % 2 == 0 and len(messages) == 2 * count:
        blocks = range(0, len(messages), 4)
        if all(messages[i : i + 2] == messages[i + 2 : i + 4] for i in blocks):
            kept = [i + offset for i in blocks for offset in (0, 1)]
            return [messages[i] for i in kept], kept
    runs: list[list[list[int]]] = []  # runs of identical adjacent exchanges
    for group in _exchanges(messages):
        key = [(messages[i]['role'], _same(messages[i]['content'])) for i in group]
        previous = runs[-1][0] if runs else None
        if previous is not None and [(messages[i]['role'], _same(messages[i]['content'])) for i in previous] == key:
            runs[-1].append(group)
        else:
            runs.append([group])
    found, reading = _readings([len(run[0]) for run in runs], [len(run) for run in runs], count)
    if not found:
        raise Unsettled(f'в тексте {len(messages)} сообщений, по «Порядку» {count}: повторы не сводятся к нему')
    if found > 1:
        raise Unsettled('повторы обменов допускают несколько прочтений при том же числе сообщений')
    kept = [i for run, multiplicity in zip(runs, reading, strict=True) for group in run[:multiplicity] for i in group]
    return [messages[index] for index in kept], kept


def _listed(value: object) -> list[str]:
    """A list column of the export ("['WEB']"); a plain value becomes a one-item list."""
    if value is None:
        return []
    text = str(value).strip()
    if text.startswith('['):
        try:
            parsed = ast.literal_eval(text)
        except (ValueError, SyntaxError):
            return [text]
        return [str(item) for item in parsed] if isinstance(parsed, list | tuple) else [str(parsed)]
    return [text] if text else []


def _meta(row: tuple, column: dict[str, int], number: int) -> dict:
    """The service columns of one row, the customer's id only as a pseudonymous key."""

    def cell(name: str) -> object:
        return _cell(row, column[name]) if name in column else None

    meta: dict = {'row': number}
    for name, key in COLUMNS.items():
        value = cell(name)
        if key in ('agents', 'actions'):
            meta[key] = _listed(value)
        elif key in ('channel', 'entrySource', 'surface'):
            meta[key] = ', '.join(_listed(value)) or None
        elif key == 'operator':
            meta[key] = None if value is None else str(value).strip().lower() == 'true'
        else:
            meta[key] = None if value is None else str(value).strip()
    # Each status code with the agents the export lists under it: which agent is under test is the profile's.
    meta['statuses'] = {
        name.removeprefix('Статус код '): _listed(cell(name)) for name in STATUS_COLUMNS if _listed(cell(name))
    }
    client = next((str(cell(name)) for name in CLIENT_COLUMNS if cell(name) not in (None, '')), None)
    meta['clientKey'] = hashlib.sha256(client.encode()).hexdigest()[:12] if client else None
    return meta


def _check_parts(data: bytes) -> None:
    """Every part of a workbook unpacks to the size it declares, all of them to at most INFLATED. openpyxl reads some
    parts whole, and a zip bomb that declares little would fill memory there; here each part is unpacked a piece at a
    time, one byte past its declared size, which a lying part has."""
    with ZipFile(io.BytesIO(data)) as archive:
        parts = archive.infolist()
        if sum(info.file_size for info in parts) > INFLATED:
            raise ValueError(
                f'Файл Excel больше {INFLATED // 1_000_000}\u00a0МБ после распаковки. '
                'Выгрузите разговоры за меньший срок.'
            )
        for info in parts:
            if info.flag_bits & 1 or info.compress_type not in (ZIP_STORED, ZIP_DEFLATED):
                raise BadZipFile('encrypted or unusual compression')  # Excel deflates; others have no bound per piece
            declared = info.file_size
            info.file_size += 1  # this archive object is only for the check
            unpacked = 0
            with archive.open(info) as part:
                while piece := part.read(1 << 16):
                    unpacked += len(piece)
            if unpacked != declared:
                raise BadZipFile('a part is not the size it declares')


def _cell(row: tuple, index: int) -> object:
    """A sheet that does not declare its size gives each row only as far as its last filled cell."""
    return row[index] if index < len(row) else None


def from_excel(data: bytes) -> list[dict]:
    """The conversations of a workbook, each with the export's service columns and how its text was read."""
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET] if SHEET in workbook.sheetnames else workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = [str(value or '').strip() for value in next(rows, ())]
        missing = [name for name in (ID, TEXT, ORDER) if name not in header]
        if missing:
            raise ValueError('В выгрузке нет колонок ' + ', '.join(f'«{name}»' for name in missing) + '.')
        column = {name: index for index, name in enumerate(header) if name}
        dialogues = []
        for number, row in enumerate(rows, 2):
            if not any(value is not None for value in row):
                continue
            layout, text_messages, messages, kept, unsettled = _read_turns(
                str(_cell(row, column[TEXT]) or ''), _cell(row, column[ORDER])
            )
            meta = _meta(row, column, number)
            status = 'as_written' if unsettled else 'exact' if len(messages) == len(text_messages) else 'collapsed'
            meta['import'] = {'status': status, 'textMessages': len(text_messages), 'layout': layout}
            if status == 'collapsed':
                meta['import']['kept'] = kept
            if unsettled:
                meta['import']['reason'] = unsettled
            dialogues.append({'id': _cell(row, column[ID]), 'messages': messages, 'meta': meta})
        return dialogues
    finally:
        workbook.close()


def from_jsonl(data: bytes) -> list[dict]:
    try:
        text = data.decode('utf-8-sig')  # Windows editors save UTF-8 with a byte order mark
    except UnicodeDecodeError as error:
        raise ValueError('Файл .jsonl не в кодировке UTF-8. Сохраните выгрузку в UTF-8.') from error
    dialogues = []
    for number, line in enumerate(text.splitlines(), 1):
        if not line.strip():
            continue
        try:
            dialogues.append(json.loads(line))
        except json.JSONDecodeError as error:
            raise ValueError(
                f'Строка {number} не читается как JSON. Нужна выгрузка чата, по одному разговору в строке.'
            ) from error
    return dialogues


def _validated(dialogues: list[dict]) -> list[dict]:
    """The conversations a check can read: the customer writes first and the agent answers. The others are left out
    (read_export counts them); an export without one usable conversation is refused."""
    usable, seen = [], set()
    for index, dialogue in enumerate(dialogues, 1):
        if not isinstance(dialogue, dict):
            raise ValueError(f'В строке {index} не объект JSON. Нужен один разговор в строке.')
        dialogue_id = dialogue.get('id')
        if isinstance(dialogue_id, bool) or not isinstance(dialogue_id, str | int) or not str(dialogue_id).strip():
            raise ValueError(f'В строке {index} у разговора нет id.')
        dialogue_id = str(dialogue_id).strip()
        if dialogue_id in seen:
            raise ValueError(f'Диалог {dialogue_id} встречается дважды.')
        seen.add(dialogue_id)
        messages = dialogue.get('messages')
        if not isinstance(messages, list):
            raise ValueError(f'В диалоге {dialogue_id} сообщения записаны не списком.')
        normalized = []
        for message in messages:
            if not isinstance(message, dict) or message.get('role') not in ('user', 'assistant'):
                raise ValueError(f'В диалоге {dialogue_id} есть сообщение с неизвестной ролью.')
            content = message.get('content')
            if not isinstance(content, str) or not content.strip():
                raise ValueError(f'В диалоге {dialogue_id} есть сообщение без текста.')
            normalized.append({'role': message['role'], 'content': content.strip()})
        if normalized and normalized[0]['role'] == 'user' and any(m['role'] == 'assistant' for m in normalized):
            kept = {'id': dialogue_id, 'messages': normalized}
            if isinstance(dialogue.get('meta'), dict):
                kept['meta'] = dialogue['meta']
            usable.append(kept)
    if not usable:
        raise ValueError('В файле нет разговоров, где клиент пишет первым и агент отвечает.')
    return usable


def prepare(name: str, data: bytes) -> list[dict]:
    """Validate an entire upload before its caller atomically replaces the stored conversations."""
    return read_export(name, data)[0]


def read_export(name: str, data: bytes) -> tuple[list[dict], int]:
    """The usable conversations of an upload, and how many it had that a check cannot read (the agent wrote first, or
    never answered): the person is told how many were left out."""
    if name.lower().endswith('.jsonl'):
        dialogues = from_jsonl(data)
    elif name.lower().endswith('.json'):
        dialogues = export_formats.from_json(data)
    elif name.lower().endswith('.csv'):
        dialogues = export_formats.from_csv(data, turns)
    elif name.lower().endswith('.xlsx'):
        try:
            _check_parts(data)
            dialogues = from_excel(data)
        except UNREADABLE as error:
            raise ValueError('Файл .xlsx повреждён или зашифрован. Сохраните выгрузку заново.') from error
    else:
        raise ValueError('Нужна выгрузка в .xlsx, .jsonl, .json или .csv.')
    usable = _validated(dialogues)
    return usable, len(dialogues) - len(usable)

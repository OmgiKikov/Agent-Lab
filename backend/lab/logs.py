"""Import and read complete recorded conversations.

Excel exports contain CLIENT/AGENT turns and the real message count in the order column. Some exports
repeat every exchange twice. Only that complete, count-confirmed export pattern is removed; a customer's
actual repeated question is preserved. Parsing is pure so a cancelled import cannot commit from a thread.
"""

import io
import json
import re
import zlib
from xml.etree.ElementTree import ParseError
from zipfile import ZIP_DEFLATED, ZIP_STORED, BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from . import store

FILE = 'logs.json'
META = 'logs-meta.json'  # the name and time of the last upload
LIMIT = 50_000_000  # an uploaded export
INFLATED = 500_000_000  # the parts of a workbook, unpacked together
SHEET = 'Данные'
ID, TEXT, ORDER = 'Id диалога', 'Текст', 'Порядок сообщения в диалоге'
# The export starts every turn on its own line; «HOST AGENT NOT FOUND» inside a message is the customer's words.
MARKER = re.compile(r'^[ \t]*(CLIENT|AGENT)\b', re.M)
# A chat button the agent sent, written into the export's text as «` ` ` transition-code CODE ` ` `».
CONTROL = re.compile(r'`\s*`\s*`\s*transition-code\s*([A-Za-z0-9_-]*)\s*`\s*`\s*`')
# The line as_seen puts under a reply for those buttons.
BUTTONS = re.compile(r'\n\[Кнопки: [^\n]*\]\Z')
# A broken workbook: openpyxl names a part the archive does not have (KeyError), zipfile meets a broken stream or a
# feature it does not read.
UNREADABLE = (BadZipFile, InvalidFileException, ParseError, KeyError, zlib.error, EOFError, NotImplementedError)


def as_seen(text: str) -> str:
    """An agent reply as the customer saw it: its words, then the buttons it sent, not the export's control code.
    A judge shown the raw code reads it as the agent's formatting (frontend/src/product/text.ts does the same)."""
    buttons = [code or 'кнопка' for code in CONTROL.findall(text)]
    words = CONTROL.sub('', text).strip()
    return words + ('\n[Кнопки: ' + ' | '.join(buttons) + ']' if buttons else '')


def words(text: str) -> str:
    """What the agent wrote in a reply, raw or as_seen: the line of buttons is the Lab's and the code is the export's,
    so neither is evidence of the agent's words."""
    return BUTTONS.sub('', CONTROL.sub('', text)).strip()


def load() -> list[dict]:
    return store.load(FILE, []) or []


def read(dialogue_id: str) -> dict | None:
    """The complete transcript, addressed by its stable imported ID."""
    return next((dialogue for dialogue in load() if dialogue['id'] == dialogue_id), None)


def turns(text: str) -> list[dict]:
    marks = list(MARKER.finditer(text))
    return [
        {
            'role': 'user' if mark.group(1) == 'CLIENT' else 'assistant',
            'content': text[mark.end() : marks[i + 1].start() if i + 1 < len(marks) else len(text)].strip(),
        }
        for i, mark in enumerate(marks)
    ]


def _message_count(order: object) -> int:
    """How many messages the order column lists. A refusal is the end of a sentence about one dialogue (from_excel)."""
    try:
        value = json.loads(str(order))
    except (ValueError, TypeError) as error:
        raise ValueError('не читается порядок сообщений') from error
    if not isinstance(value, list) or not value:
        raise ValueError('пустой порядок сообщений')
    return len(value)


def _export_messages(messages: list[dict], count: int) -> list[dict]:
    if len(messages) == count:
        return messages
    # Known duplicated-export layout: customer+agent, the same customer+agent, then the next exchange.
    # Count is essential: equality of adjacent exchanges alone cannot distinguish an export from a real repeat.
    if count % 2 == 0 and len(messages) == 2 * count:
        blocks = [messages[index : index + 4] for index in range(0, len(messages), 4)]
        if all(block[:2] == block[2:] for block in blocks):
            return [message for block in blocks for message in block[:2]]
    raise ValueError('текст не совпадает с порядком сообщений')


def _check_parts(data: bytes) -> None:
    """Every part of a workbook unpacks to the size it declares, all of them to at most INFLATED. openpyxl reads some
    parts whole, and a zip bomb that declares little would fill memory there; here each part is unpacked a piece at a
    time, one byte past its declared size, which a lying part has."""
    with ZipFile(io.BytesIO(data)) as archive:
        parts = archive.infolist()
        if sum(info.file_size for info in parts) > INFLATED:
            raise ValueError(
                f'Файл Excel больше {INFLATED // 1_000_000}\u00a0МБ после распаковки. '
                'Выгрузите разговоры за меньший срок.'
            )
        for info in parts:
            if info.flag_bits & 1 or info.compress_type not in (ZIP_STORED, ZIP_DEFLATED):
                raise BadZipFile('encrypted or unusual compression')  # Excel deflates; others have no bound per piece
            declared = info.file_size
            info.file_size += 1  # this archive object is only for the check
            unpacked = 0
            with archive.open(info) as part:
                while piece := part.read(1 << 16):
                    unpacked += len(piece)
            if unpacked != declared:
                raise BadZipFile('a part is not the size it declares')


def _cell(row: tuple, index: int) -> object:
    """A sheet that does not declare its size gives each row only as far as its last filled cell."""
    return row[index] if index < len(row) else None


def from_excel(data: bytes) -> list[dict]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET] if SHEET in workbook.sheetnames else workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = [str(value or '').strip() for value in next(rows, ())]
        missing = [name for name in (ID, TEXT, ORDER) if name not in header]
        if missing:
            raise ValueError('В выгрузке нет колонок ' + ', '.join(f'«{name}»' for name in missing) + '.')
        column = {name: header.index(name) for name in (ID, TEXT, ORDER)}
        dialogues = []
        for row in rows:
            if not any(value is not None for value in row):
                continue
            dialogue_id = _cell(row, column[ID])
            try:
                count = _message_count(_cell(row, column[ORDER]))
                messages = _export_messages(turns(str(_cell(row, column[TEXT]) or '')), count)
            except ValueError as error:
                raise ValueError(f'В диалоге {dialogue_id} {error}.') from error
            dialogues.append({'id': dialogue_id, 'messages': messages})
        return dialogues
    finally:
        workbook.close()


def from_jsonl(data: bytes) -> list[dict]:
    try:
        text = data.decode('utf-8-sig')  # Windows editors save UTF-8 with a byte order mark
    except UnicodeDecodeError as error:
        raise ValueError('Файл .jsonl не в кодировке UTF-8. Сохраните выгрузку в UTF-8.') from error
    dialogues = []
    for number, line in enumerate(text.splitlines(), 1):
        if not line.strip():
            continue
        try:
            dialogues.append(json.loads(line))
        except json.JSONDecodeError as error:
            raise ValueError(
                f'Строка {number} не читается как JSON. Нужна выгрузка чата, по одному разговору в строке.'
            ) from error
    return dialogues


def _validated(dialogues: list[dict]) -> list[dict]:
    """The conversations a check can read: the customer writes first and the agent answers. The others are left out
    (read_export counts them); an export without one usable conversation is refused."""
    usable, seen = [], set()
    for index, dialogue in enumerate(dialogues, 1):
        if not isinstance(dialogue, dict):
            raise ValueError(f'В строке {index} не объект JSON. Нужен один разговор в строке.')
        dialogue_id = dialogue.get('id')
        if isinstance(dialogue_id, bool) or not isinstance(dialogue_id, str | int) or not str(dialogue_id).strip():
            raise ValueError(f'В строке {index} у разговора нет id.')
        dialogue_id = str(dialogue_id).strip()
        if dialogue_id in seen:
            raise ValueError(f'Диалог {dialogue_id} встречается дважды.')
        seen.add(dialogue_id)
        messages = dialogue.get('messages')
        if not isinstance(messages, list):
            raise ValueError(f'В диалоге {dialogue_id} сообщения записаны не списком.')
        normalized = []
        for message in messages:
            if not isinstance(message, dict) or message.get('role') not in ('user', 'assistant'):
                raise ValueError(f'В диалоге {dialogue_id} есть сообщение с неизвестной ролью.')
            content = message.get('content')
            if not isinstance(content, str) or not content.strip():
                raise ValueError(f'В диалоге {dialogue_id} есть сообщение без текста.')
            normalized.append({'role': message['role'], 'content': content.strip()})
        if normalized and normalized[0]['role'] == 'user' and any(m['role'] == 'assistant' for m in normalized):
            usable.append({'id': dialogue_id, 'messages': normalized})
    if not usable:
        raise ValueError('В файле нет разговоров, где клиент пишет первым и агент отвечает.')
    return usable


def prepare(name: str, data: bytes) -> list[dict]:
    """Validate an entire upload before its caller atomically replaces the stored conversations."""
    return read_export(name, data)[0]


def read_export(name: str, data: bytes) -> tuple[list[dict], int]:
    """The usable conversations of an upload, and how many it had that a check cannot read (the agent wrote first, or
    never answered): the person is told how many were left out."""
    if name.lower().endswith('.jsonl'):
        dialogues = from_jsonl(data)
    elif name.lower().endswith('.xlsx'):
        try:
            _check_parts(data)
            dialogues = from_excel(data)
        except UNREADABLE as error:
            raise ValueError('Файл .xlsx повреждён или зашифрован. Сохраните выгрузку заново.') from error
    else:
        raise ValueError('Нужна выгрузка в .xlsx или .jsonl.')
    usable = _validated(dialogues)
    return usable, len(dialogues) - len(usable)


def commit(dialogues: list[dict], name: str | None = None) -> int:
    """The new export with its file name, in one transaction: a saved check never names the previous file."""
    store.replace_inputs(FILE, dialogues, {META: {'file': name, 'updatedAt': store.now()}} if name else None)
    return len(dialogues)


def meta() -> dict:
    """The export the conversations came from: its file name and when it was uploaded, when that is known."""
    return store.load(META) or {'file': None, 'updatedAt': None}
